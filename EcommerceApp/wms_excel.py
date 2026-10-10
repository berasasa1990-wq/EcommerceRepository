"""Validated WMS product import; one row per product/location."""
from decimal import Decimal, InvalidOperation
from io import BytesIO
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .models import Product, WMSLocation, ProductWMSStock, Category, Brand, Tag

HEADERS = ['Naziv', 'Šifra', 'Cijena', 'Akcijska cijena', 'Barkod', 'Lokacija', 'Količina', 'Opis', 'Kategorija', 'Brend', 'Tagovi']


def template_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Artikli'
    sheet.append(HEADERS)
    sheet.append(['Testni artikal', 'TEST-001', 100, 90, '0012345678901', 'A-10', 5, 'Detaljan opis testnog artikla.', 'Oprema > Dodaci', 'Testni brend', 'novo, dodatak'])
    sheet.append(['Testni artikal', 'TEST-001', 100, 90, '0012345678901', 'A-12', 2, 'Detaljan opis testnog artikla.', 'Oprema > Dodaci', 'Testni brend', 'novo, dodatak'])
    sheet.append(['Drugi testni artikal', 'TEST-002', 50, None, '0012345678902', None, None, 'Opis drugog artikla.', 'Oprema > Ostalo', '', ''])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='008FC3')
        cell.alignment = Alignment(horizontal='center')
    for col, width in zip('ABCDEFGHIJK', [32, 22, 16, 20, 24, 22, 16, 55, 40, 24, 30]):
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows(min_row=2):
        for index in (1, 4):
            row[index].number_format = '@'
        for index in (2, 3):
            row[index].number_format = '0.00'
    guide = book.create_sheet('Uputstvo')
    for text in [
        'Obrišite testne redove i unesite svoje artikle u list Artikli.',
        'Naziv, Šifra i Cijena su obavezni. Šifra prepoznaje postojeći artikal.',
        'Šifru i barkod unesite kao tekst da sačuvate početne nule.',
        'Akcijska cijena je opciona. Prazno polje uklanja akcijsku cijenu.',
        'Za više lokacija ponovite iste podatke artikla, uz drugu lokaciju i količinu.',
        'Lokacija i Količina unose se zajedno. Količina je cijeli broj 0 ili veći.',
        'Količina POSTAVLJA stanje te lokacije; ne dodaje na postojeće stanje.',
        'Lokacije koje nisu navedene ostaju nepromijenjene. Nova lokacija se kreira.',
        'Prazna lokacija i količina ne mijenjaju postojeće zalihe.',
        'Opis je tekst. Prazna dodatna polja ne mijenjaju postojeće podatke.',
        'Kategorija je puna putanja, npr. Ribolov > Štapovi > Šaranski štapovi.',
        'Nove kategorije i podkategorije automatski se kreiraju i povezuju.',
        'Brend je naziv brenda. Tagovi se odvajaju zarezom.',
        'Import provjerava sve redove. Ako ima greške, ništa se ne upisuje.'
    ]:
        guide.append([text])
    guide.column_dimensions['A'].width = 110
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def import_products(upload):
    book = load_workbook(upload, read_only=True, data_only=False)
    try:
        sheet = book['Artikli'] if 'Artikli' in book.sheetnames else book.active
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows, ()))
        if headers not in (HEADERS, HEADERS[:7]):
            raise ValueError('Kolone moraju odgovarati preuzetom templateu.')
        products = {}
        def text(value):
            return str(value).strip() if value is not None else ''
        def price(value, optional=False):
            if optional and value in (None, ''):
                return None
            result = Decimal(text(value).replace(',', '.'))
            if not result.is_finite() or result < 0 or result > Decimal('99999999.99'):
                raise ValueError
            return result.quantize(Decimal('0.01'))
        for number, row in enumerate(rows, 2):
            if number > 10001:
                raise ValueError('Maksimalno 10000 redova po importu.')
            if not any(value is not None for value in row):
                continue
            row = list(row[:len(HEADERS)]) + [None] * max(0, len(HEADERS) - len(row))
            try:
                name, code, barcode, location = text(row[0]), text(row[1]), text(row[4]), text(row[5])
                description, category_path, brand, tags = map(text, row[7:11])
                category_parts = [part.strip() for part in category_path.split('>')] if category_path else []
                if any(not part or len(part) > 100 for part in category_parts) or len(brand) > Brand._meta.get_field('naziv').max_length or any(len(tag.strip()) > 50 for tag in tags.split(',')):
                    raise ValueError
                data = (name, price(row[2]), price(row[3], True), barcode, description, category_path, brand, tags)
                if not name or not code or len(name) > 200 or any(text(v).startswith('=') for v in row):
                    raise ValueError
                Product._meta.get_field('sifra').clean(code, None)
                Product._meta.get_field('barkod').clean(barcode, None)
                if bool(location) != (row[6] not in (None, '')) or len(location) > 120:
                    raise ValueError
                quantity = None
                if location:
                    qty = Decimal(text(row[6]))
                    if not qty.is_finite() or qty < 0 or qty != int(qty) or qty > 2147483647:
                        raise ValueError
                    quantity = int(qty)
                entry = products.setdefault(code, {'data': data, 'locations': {}})
                if entry['data'] != data or (location and location in entry['locations']):
                    raise ValueError
                if location:
                    entry['locations'][location] = quantity
            except Exception as exc:
                raise ValueError(f'Red {number}: provjerite obavezne podatke, cijene, lokaciju i količinu. Podaci za istu šifru moraju biti isti; lokacija se ne ponavlja.') from exc
        if not products:
            raise ValueError('Datoteka nema artikala za import.')
        with transaction.atomic():
            for code, entry in products.items():
                matches = list(Product.objects.select_for_update().filter(sifra=code))
                if len(matches) > 1:
                    raise ValueError(f'Šifra {code} postoji na više artikala. Ispravite duplikate.')
                product = matches[0] if matches else Product(sifra=code, stanje=0)
                product.naziv, product.cijena, product.akcijska_cijena, product.barkod = entry['data'][:4]
                description, category_path, brand, tags = entry['data'][4:]
                if description:
                    product.opis = description
                if category_path:
                    parent = None
                    for part in category_path.split('>'):
                        matches = list(Category.objects.filter(naziv=part.strip(), roditelj=parent))
                        if len(matches) > 1:
                            raise ValueError(f'Kategorija {part.strip()} ima duplikate.')
                        parent = matches[0] if matches else Category.objects.create(naziv=part.strip(), roditelj=parent)
                    product.kategorija = parent
                if brand:
                    product.brend, _ = Brand.objects.get_or_create(naziv=brand)
                product.save()
                if tags:
                    product.tagovi.set([Tag.get_or_create_by_name(tag.strip())[0] for tag in tags.split(',') if tag.strip()])
                for name, quantity in entry['locations'].items():
                    location, _ = WMSLocation.objects.get_or_create(naziv=name)
                    ProductWMSStock.objects.update_or_create(product=product, lokacija=location, defaults={'kolicina': quantity})
                if entry['locations']:
                    product.stanje = sum(product.wms_zalihe.values_list('kolicina', flat=True))
                    if product.stanje > 2147483647:
                        raise ValueError(f'Prevelika ukupna količina za šifru {code}.')
                    ids = list(product.wms_zalihe.values_list('lokacija_id', flat=True))
                    product.wms_lokacija_id = ids[0] if len(ids) == 1 else None
                    product.save(update_fields=['stanje', 'wms_lokacija', 'na_stanju'])
        return len(products)
    finally:
        book.close()
